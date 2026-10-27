"""Equivalent factual report content rendered locally as PDF and DOCX."""
from io import BytesIO
from pathlib import Path
import json
from xml.sax.saxutils import escape

DISCLAIMER = (
    "Decision support for human review only. This report does not establish legal correctness, "
    "enforceability or completeness. Anomaly scores are not calibrated probabilities of legal error. "
    "Intervals use synthetic-shuffle calibration and do not quantify legal certainty. "
    "Rule matches are automated textual candidates, not expert-verified legal defects. "
    "Research benchmarks remain exploratory; independent annotation and unseen-family evaluation are pending."
)


def report_blocks(result: dict) -> list[tuple[str, str]]:
    meta = result.get("metadata", {})
    scoring = result.get("scoring_result", {})
    clauses = scoring.get("clauses", [])
    inventory = {c["clause_id"]: c for c in meta.get("clause_inventory", [])}
    blocks = [("title", "LexAudit Contract Analysis Report"),
              ("text", f"Document: {meta.get('document_name', result['doc_id'])}"),
              ("text", f"Document ID: {result['doc_id']}"),
              ("text", f"Analysis date: {meta.get('timestamp', 'Unavailable')}"),
              ("text", f"Model/version: {meta.get('model_version', 'Local dual-channel auditor v1')}"),
              ("heading", "Executive summary"),
              ("text", f"{result['total_clauses']} clauses analyzed; {result['ml_anomaly_count']} statistical anomalies; "
               f"{result['rule_violation_count']} automated textual rule candidates. "
               f"Mean anomaly score: {scoring.get('mean_combined_score', 0)}. "
               f"Review state: {meta.get('analysis_state', 'REVIEW_REQUIRED' if result['total_flags'] else 'NO_FINDINGS')}. "
               "No findings means no configured detector flags, not a clean legal bill of health."),
              ("heading", "Score and confidence interpretation"), ("text", DISCLAIMER),
              ("heading", "Textual rule candidates")]
    if not result.get("rule_flags"):
        blocks.append(("text", "No textual rule candidates raised."))
    for flag in sorted(result.get("rule_flags", []), key=lambda f: {"HIGH": 0, "MEDIUM": 1}.get(f["severity"], 2)):
        blocks.extend([("subheading", f"{flag['severity']} - {flag['title']}"),
                       ("text", flag["description"]),
                       ("text", f"Clause evidence: {flag.get('clause_id')}; involved: {', '.join(flag.get('involved_clause_ids', []))}"),
                       ("text", "Evidence: " + json.dumps(flag.get("evidence", {}), ensure_ascii=False)),
                       ("text", "Claim scope: automated textual check; human review required.")])
    structure = meta.get("structural_document")
    if structure:
        blocks.extend([("heading", "Source-preserving structural evidence"),
                       ("text", f"Parser: {structure.get('version')}; offsets: {structure.get('coordinate_space')}. "
                        "Unicode code points in extracted text, not PDF page coordinates. Neural scoring uses its unchanged legacy inventory."),
                       ("text", f"Source SHA-256: {structure.get('source_sha256')}.")])
        for section in structure.get("sections", []):
            blocks.append(("text", f"{section['target_type']} {section['identity']}; parent: {section.get('parent_identity')}; "
                           f"source offsets {section['heading_start']} to {section['heading_end']}: {section['heading_text']}"))
        for review in meta.get("reference_reviews", []):
            blocks.append(("text", f"REVIEW_REQUIRED: {review.get('text', review.get('heading_text', ''))}. "
                           f"{review.get('reason')}; offsets {review.get('char_start', review.get('heading_start'))} "
                           f"to {review.get('char_end', review.get('heading_end'))}."))
    temporal = meta.get('temporal_analysis')
    if temporal:
        blocks.extend([('heading','Temporal analysis'),
            ('text',f"State: {temporal.get('status')}; version: {temporal.get('version')}; coordinates: {temporal.get('coordinate_space')}. "
             'Confirmed means a supported deterministic textual contradiction. Potential differences, uncertainty and unsupported syntax are separate; severity is review priority, not legal certainty.')])
        for statement in temporal.get('statements',[]):
            blocks.append(('text',f"{statement['statement_id']} — scope {statement.get('scope')}; phase {statement.get('phase')}; "
                f"actor {statement.get('actor') or 'not stated'}; action {statement.get('action') or 'not stated'}; "
                f"trigger {statement.get('triggering_event') or 'not stated'}; date {statement.get('absolute_date')}; "
                f"duration {statement.get('duration_value')} {statement.get('duration_unit')}; basis {statement.get('timing_basis')}; "
                f"operator {statement.get('deadline_type')}; qualifiers {statement.get('qualifiers')}; "
                f"source {statement['char_start']} to {statement['char_end']}: {statement['raw_text']}"))
        for decision in temporal.get('decisions',[]):
            blocks.append(('text',f"{decision['category']} / {decision.get('check', 'bounded parse')}: {decision['reason']} "
                f"Statements: {', '.join(decision.get('statement_ids', []))}."))
    blocks.append(("heading", "Statistical anomalies"))
    anomalies = [c for c in clauses if c["is_anomaly"]]
    if not anomalies:
        blocks.append(("text", "No statistical anomalies raised."))
    for c in sorted(anomalies, key=lambda c: -c["combined_score"]):
        blocks.extend([("subheading", f"{c['severity']} - {c['clause_id']}"),
                       ("keeptext", f"Anomaly score: {c['combined_score']}; Channel A: {c['channel_a_score']}; Channel B: {c['channel_b_score']}."),
                       ("keeptext", c.get("severity_justification") or "Raised by the frozen ensemble thresholds; requires review."),
                       ("text", f"Synthetic interval: {c.get('confidence_interval')}; agreement: {c.get('agreement_type')}.")])
    blocks.append(("heading", "Clause inventory and locations"))
    if not clauses:
        blocks.append(("text", "No clauses available in this analysis."))
    for c in clauses:
        source = inventory.get(c["clause_id"], {})
        blocks.extend([("subheading", f"Clause {c['sequence_idx'] + 1} - {c['clause_id']}"),
                       ("text", f"Normalized text offsets: {source.get('char_start', 'unavailable')} to {source.get('char_end', 'unavailable')}. "
                        f"Severity: {c['severity']}; score: {c['combined_score']}."),
                       ("text", source.get("text", c["text_preview"]))])
    blocks.append(("heading", "Experimental dependency evidence"))
    graph = meta.get("dependency_graph")
    if graph:
        blocks.append(("text", f"State: {graph.get('status')}. {graph.get('scope', '')} {graph.get('reason', '')}"))
        for finding in [*graph.get("findings", []), *graph.get("reviews", [])]:
            target = finding.get("target", {})
            blocks.extend([("subheading", f"{finding.get('status')} - {finding.get('source_clause_id')}"),
                           ("text", f"Reference: {finding.get('reference')}. {finding.get('reason')}. "
                            f"Target section: {target.get('section_id', 'unresolved')}; clause: {target.get('clause_id', 'unresolved')}. "
                            f"Reference offsets: {finding.get('source_char_start', finding.get('char_start', 'unavailable'))} "
                            f"to {finding.get('source_char_end', finding.get('char_end', 'unavailable'))}; "
                            f"coordinate space: {finding.get('coordinate_space', 'legacy source clause')}. "
                            f"Clause association: {finding.get('clause_offset_mapping', 'legacy')}.")])
        if not graph.get("findings") and not graph.get("reviews"):
            blocks.append(("text", "No explicit directional conflicts or unresolved targets found."))
    else:
        blocks.append(("text", "Opt-in dependency graph has not been run for this analysis."))
    blocks.append(("heading", "Counterfactual findings"))
    counterfactual = meta.get("counterfactual")
    if counterfactual:
        blocks.append(("text", f"Target: {counterfactual.get('target_clause_id')}. "
                       f"Original mean score: {counterfactual.get('original_mean_combined_score')}. "
                       f"Best objective delta: {counterfactual.get('best_objective_delta')}. "
                       f"Abstaining: {counterfactual.get('is_abstaining')}. "
                       f"Reason: {counterfactual.get('abstention_reason') or 'See candidate evidence'}. "
                       "These are unverified suggestions, never automatic repairs."))
        for candidate in counterfactual.get("candidates", []):
            blocks.append(("text", f"Position {candidate['position_idx']}: mean combined score {candidate['mean_combined_score']}; "
                           f"objective delta {candidate['objective_delta']}; recommended {candidate['is_recommended']}."))
    else:
        blocks.append(("text", "No counterfactual evaluation saved for this analysis. Suggestions are unverified and never applied automatically."))
    blocks.extend([("heading", "Limitations and human review"), ("text", DISCLAIMER),
                   ("heading", "Technical appendix"),
                   ("text", f"Checkpoint SHA-256: {meta.get('checkpoint_sha256', 'unavailable')}. "
                    f"Calibration source: {scoring.get('calibration_source', 'unavailable')}. "
                    "Locations refer to normalized extracted text, not PDF page coordinates. "
                    "The PDF and DOCX share the same source results. Unsupported font glyphs are represented as Unicode code points."),
                   ("text", "Scoring configuration: " + json.dumps(scoring.get("metadata", {}), ensure_ascii=False))])
    return blocks


def render_report(result: dict, format: str) -> bytes:
    blocks = report_blocks(result)
    output = BytesIO()
    if format == "docx":
        from docx import Document
        from docx.shared import Inches, Pt
        from docx.shared import RGBColor
        from docx.oxml.ns import qn
        doc = Document()
        section = doc.sections[0]
        section.top_margin = section.bottom_margin = Inches(0.75)
        section.left_margin = section.right_margin = Inches(0.85)
        doc.styles["Normal"].font.name = "Arial"
        doc.styles["Normal"].font.size = Pt(10)
        doc.styles["Title"].font.color.rgb = RGBColor(0, 0, 0)
        for style in doc.styles:
            if style.type == 1:
                properties = style.element.find(qn("w:pPr"))
                if properties is not None:
                    border = properties.find(qn("w:pBdr"))
                    if border is not None:
                        properties.remove(border)
        for kind, text in blocks:
            # XML cannot contain control characters, even when PDF rendering tolerates them.
            text = "".join(c if c >= " " or c in "\n\t" else f"[U+{ord(c):04X}]" for c in text)
            if kind == "title":
                doc.add_paragraph(text, "Title")
            elif kind in {"heading", "subheading"}:
                doc.add_heading(text, 1 if kind == "heading" else 2)
            else:
                paragraph = doc.add_paragraph(text)
                if kind == "keeptext":
                    paragraph.paragraph_format.keep_with_next = True
        section.footer.paragraphs[0].text = "LexAudit | Local analysis | Human review required"
        doc.save(output)
    elif format == "pdf":
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_LEFT
        font_paths = [Path(__file__).parents[2] / "frontend/public/fonts/DejaVuSans.ttf",
                      Path("C:/Windows/Fonts/arial.ttf"), Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf")]
        font = next((p for p in font_paths if p.exists()), None)
        if font is None:
            raise RuntimeError("Install Arial or DejaVu Sans for Unicode PDF export.")
        if "LexAudit" not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont("LexAudit", str(font)))
        supported = pdfmetrics.getFont("LexAudit").face.charToGlyph
        styles = getSampleStyleSheet()
        for style in styles.byName.values():
            style.fontName = "LexAudit"
            style.alignment = TA_LEFT
            style.wordWrap = "CJK"
        styles["BodyText"].fontSize = 9
        styles["BodyText"].leading = 13
        styles["Heading1"].textColor = colors.HexColor("#173b55")
        for heading in ("Title", "Heading1", "Heading2"):
            styles[heading].keepWithNext = True
        styles.add(ParagraphStyle("KeepBody", parent=styles["BodyText"], keepWithNext=True))
        style_names = {"title": "Title", "heading": "Heading1", "subheading": "Heading2", "text": "BodyText", "keeptext": "KeepBody"}
        story = []
        for kind, text in blocks:
            text = "".join(c if ord(c) in supported or c in "\n\t" else f"[U+{ord(c):04X}]" for c in text)
            spacing = Spacer(1, 6)
            # A spacer must not break a heading/explanation keep-with-next chain.
            spacing.keepWithNext = True
            story.extend([Paragraph(escape(text).replace("\n", "<br/>"), styles[style_names[kind]]), spacing])
        def footer(canvas, document):
            canvas.setFont("LexAudit", 8)
            canvas.drawString(48, 28, "LexAudit | Human review required")
            canvas.drawRightString(document.pagesize[0] - 48, 28, str(document.page))
        SimpleDocTemplate(output, title="LexAudit Contract Analysis Report", leftMargin=48, rightMargin=48,
                          topMargin=48, bottomMargin=48).build(story, onFirstPage=footer, onLaterPages=footer)
    else:
        raise ValueError("Unsupported report format")
    return output.getvalue()
