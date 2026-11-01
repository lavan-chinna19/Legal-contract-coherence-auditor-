"""Release acceptance using synthetic text and actual local renderers/models."""
from io import BytesIO
import json
import time
import zipfile
import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from src.api.main import app
from src.api import storage, jobs
from src.api.ingestion import parse_document, MAX_UPLOAD_BYTES
from src.api.reports import render_report, report_blocks
from src.api.dependencies import get_unified_auditor, get_ml_segmenter
from src.rules.dependency_graph import detect_dependencies

SAMPLE = "Section 3. Payment\nPayment is due under Section 1 above.\nSection 1. Definitions\nServices means work described by this agreement.\nSection 2. Confidentiality\nEach party shall protect information."
KEY = {"X-API-Key": "release_test_owner"}
OTHER = {"X-API-Key": "release_test_other"}


def test_concurrent_model_initialization_constructs_once(monkeypatch):
    from concurrent.futures import ThreadPoolExecutor
    import threading
    import src.embeddings.frozen as frozen
    barrier = threading.Barrier(8)
    calls = []

    class FrozenModel:
        def parameters(self):
            return []

        def eval(self):
            return self

    def construct(*args, **kwargs):
        calls.append(1)
        time.sleep(0.02)
        return FrozenModel()

    monkeypatch.setattr(frozen, "SentenceTransformer", construct)
    embedder = frozen.FrozenLegalBERTEmbedder()

    def load():
        barrier.wait(timeout=5)
        embedder._load_model()
        return embedder._model

    with ThreadPoolExecutor(max_workers=8) as pool:
        models = list(pool.map(lambda _: load(), range(8)))
    assert len(calls) == 1
    assert all(model is models[0] for model in models)


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("API_KEYS", "release_test_owner,release_test_other")
    from src.api import rate_limiter as limits
    # Each case gets a fresh production-configured bucket, rather than carrying
    # another case's requests or disabling/throttling the actual limiter.
    monkeypatch.setattr(limits, "rate_limiter", limits.SlidingWindowRateLimiter())
    monkeypatch.setattr(storage, "UPLOAD_DIR", tmp_path)
    import src.api.retention as retention
    monkeypatch.setattr(retention, "UPLOAD_DIR", tmp_path)
    return TestClient(app)


def clause(text, index):
    return {"clause_id": f"c{index}", "text": text, "sequence_idx": index}


@pytest.mark.parametrize("direction,expected", [("above", "FINDING"), ("below", "NO_FINDINGS"), ("", "NO_FINDINGS")])
def test_directional_and_nondirectional_graph(direction, expected):
    graph = detect_dependencies([clause(f"Section 3. Payment\nSee Section 1 {direction}.", 0), clause("Section 1. Terms", 1)])
    assert graph["status"] == expected
    assert graph["edges"][0]["target"]["clause_id"] == "c1"


def test_graph_abstention_and_local_antecedent():
    for targets in [[], [clause("Section 1. Terms", 1), clause("Section 1. Reused", 2)]]:
        graph = detect_dependencies([clause("Section 3. Payment\nSee Section 1 above.", 0), *targets])
        assert graph["reviews"] and not graph["findings"]
    intact = clause("Section 1. Terms\nThe fee is ten dollars. Notwithstanding the foregoing, the fee may be waived.", 1)
    assert not detect_dependencies([intact])["findings"]
    assert not detect_dependencies([clause("Section 5. Other", 0), intact])["findings"]
    local = clause("Section 1. Terms\nSection 1 above defines the fee.", 0)
    assert detect_dependencies([local])["status"] == "NO_FINDINGS"
    assert detect_dependencies([clause("Section 9. Other", 1), local])["status"] == "NO_FINDINGS"
    assert detect_dependencies([clause("Section 3. Caveat\nNotwithstanding the foregoing, the obligation survives.", 0)])["status"] == "REVIEW_REQUIRED"
    assert detect_dependencies([clause("See Section IV above.", 0)])["status"] == "REVIEW_REQUIRED"


def test_immediate_predecessor():
    good = [clause("Section 1. Terms", 0), clause("Section 2. Fees", 1),
            clause("Section 3. Payment\nSee immediately preceding Section 2.", 2)]
    assert detect_dependencies(good)["status"] == "NO_FINDINGS"
    bad = [good[1], good[0], good[2]]
    assert detect_dependencies(bad)["status"] == "FINDING"


@pytest.mark.parametrize("name,data,status", [("a.exe", b"hello", 415), ("a.pdf", b"broken", 400),
    ("a.docx", b"broken", 400), ("a.txt", b"\xff", 400), ("a.txt", b"", 400),
    ("a.txt", b"x" * (MAX_UPLOAD_BYTES + 1), 413)], ids=["unsupported", "corrupt-pdf", "corrupt-docx", "invalid-utf8", "empty", "oversize"])
def test_invalid_uploads(client, name, data, status):
    response = client.post("/v1/upload", files={"file": (name, data)}, headers=KEY)
    assert response.status_code == status


def test_docx_parser_and_archive_limits():
    from docx import Document
    doc = Document()
    doc.add_paragraph(SAMPLE)
    stream = BytesIO(); doc.save(stream)
    assert "Section 1" in parse_document(stream.getvalue(), "sample.docx")
    malicious = BytesIO()
    with zipfile.ZipFile(malicious, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "x" * (26 * 1024 * 1024))
    with pytest.raises(HTTPException) as exc:
        parse_document(malicious.getvalue(), "bomb.docx")
    assert exc.value.status_code == 413


def empty_result():
    return {"doc_id": "doc_empty", "total_clauses": 0, "ml_anomaly_count": 0, "rule_violation_count": 0,
            "total_flags": 0, "metadata": {}, "scoring_result": {"clauses": []}, "rule_flags": []}


def extract_reports(result):
    from docx import Document
    import pdfplumber
    pdf = render_report(result, "pdf")
    docx = render_report(result, "docx")
    assert pdf.startswith(b"%PDF") and zipfile.is_zipfile(BytesIO(docx))
    with pdfplumber.open(BytesIO(pdf)) as document:
        pdf_text = "\n".join(p.extract_text() or "" for p in document.pages)
    docx_text = "\n".join(p.text for p in Document(BytesIO(docx)).paragraphs)
    return pdf, docx, pdf_text, docx_text


def test_empty_and_uncertain_reports():
    result = empty_result()
    _, _, pdf_text, docx_text = extract_reports(result)
    for text in [pdf_text, docx_text]:
        for required in ["LexAudit", "Executive summary", "Clause inventory", "No statistical anomalies", "not calibrated probabilities", "Technical appendix"]:
            assert required in text
    result["metadata"] = {"analysis_state": "REVIEW_REQUIRED", "dependency_graph": {"status": "REVIEW_REQUIRED", "reason": "ambiguous target"}}
    for text in extract_reports(result)[2:]:
        assert "REVIEW_REQUIRED" in text and "ambiguous target" in text


def test_real_analysis_saved_evidence_and_exports(client, monkeypatch):
    response = client.post("/v1/upload", files={"file": ("../../synthetic.txt", SAMPLE.encode())}, headers=KEY)
    assert response.status_code == 200
    doc_id = response.json()["doc_id"]
    assert storage.document_name(doc_id) == "synthetic.txt"
    preview_response = client.get(f"/v1/documents/{doc_id}/clauses", headers=KEY)
    assert preview_response.status_code == 200
    preview = preview_response.json()
    assert client.get(f"/v1/documents/{doc_id}/clauses", headers=OTHER).status_code == 404
    assert client.get(f"/v1/documents/{doc_id}/clauses").status_code == 401
    assert SAMPLE.encode() not in (storage.UPLOAD_DIR / f"{doc_id}.enc").read_bytes()
    response = client.post("/v1/analyze", json={"doc_id": doc_id}, headers=KEY)
    assert response.status_code == 200, response.text
    job_id = response.json()["job_id"]
    for _ in range(120):
        status = client.get(f"/v1/jobs/{job_id}", headers=KEY).json()
        if status["status"] in {"COMPLETED", "FAILED"}:
            break
        time.sleep(0.25)
    assert status["status"] == "COMPLETED", status
    result = client.get(f"/v1/jobs/{job_id}/results", headers=KEY).json()
    inventory = result["metadata"]["clause_inventory"]
    assert preview["clauses"] == inventory
    source = result["metadata"]["source_text"]
    assert {c["clause_id"] for c in inventory} == {c["clause_id"] for c in result["scoring_result"]["clauses"]}
    for c in inventory:
        assert source[c["char_start"]:c["char_end"]] == c["text"]
    before_scores = json.dumps(result["scoring_result"], sort_keys=True)
    assert client.post("/v1/experimental/dependencies", json={"job_id": job_id, "experimental": False}, headers=KEY).status_code == 400
    graph = client.post("/v1/experimental/dependencies", json={"job_id": job_id, "experimental": True}, headers=KEY)
    assert graph.status_code == 200 and graph.json()["findings"]
    for suffix in ["", "/results", "/report/pdf", "/report/docx"]:
        assert client.get(f"/v1/jobs/{job_id}{suffix}", headers=OTHER).status_code == 404
        assert client.get(f"/v1/jobs/{job_id}{suffix}").status_code == 401
    assert client.post("/v1/analyze", json={"doc_id": doc_id}, headers=OTHER).status_code == 404
    assert client.get("/v1/documents", headers=OTHER).json()["documents"] == []
    jobs.JOB_STORE.pop(job_id)
    restored = client.get(f"/v1/jobs/{job_id}/results", headers=KEY).json()
    assert json.dumps(restored["scoring_result"], sort_keys=True) == before_scores
    assert restored["metadata"]["dependency_graph"]["findings"]
    pdf, docx, pdf_text, docx_text = extract_reports(restored)
    for fmt, expected in [("pdf", pdf), ("docx", docx)]:
        response = client.get(f"/v1/jobs/{job_id}/report/{fmt}", headers=KEY)
        assert response.status_code == 200 and len(response.content) > 1000
        assert "attachment" in response.headers["content-disposition"]
        assert response.headers["cache-control"] == "no-store"
    for text in [pdf_text, docx_text]:
        assert doc_id in text and "synthetic.txt" in text and "Payment is due" in text
    assert client.get("/v1/documents", headers=KEY).json()["documents"][0]["job_id"] == job_id


def test_model_failure_and_resource_busy(client, monkeypatch):
    monkeypatch.setattr(jobs, "MODEL_WAIT_SECONDS", 0)
    from src.api.dependencies import get_auditor_dep
    def unavailable():
        raise HTTPException(503, "Analysis model unavailable")
    app.dependency_overrides[get_auditor_dep] = unavailable
    try:
        uploaded = client.post("/v1/upload", files={"file": ("preview-without-model.txt", SAMPLE.encode())}, headers=KEY)
        assert uploaded.status_code == 200
        doc_id = uploaded.json()["doc_id"]
        preview = client.get(f"/v1/documents/{doc_id}/clauses", headers=KEY)
        assert preview.status_code == 200 and preview.json()["clauses"]
        response = client.post("/v1/analyze", json={"doc_id": doc_id}, headers=KEY)
        assert response.status_code == 503
    finally:
        app.dependency_overrides.clear()
    jobs.ANALYSIS_LOCK.acquire()
    try:
        with pytest.raises(HTTPException) as exc:
            jobs.run_model_operation(lambda: None)
        assert exc.value.status_code == 429
    finally:
        jobs.ANALYSIS_LOCK.release()


def test_path_traversal_and_session(client):
    assert client.get("/v1/session", headers=KEY).status_code == 200
    assert client.get("/v1/session").status_code == 401
    assert not storage.document_exists("../../secret")
    with pytest.raises(HTTPException):
        storage.load_and_decrypt_document("../../secret")


def test_multipart_limits_before_parsing(client):
    from starlette.formparsers import MultiPartParser
    from src.api.upload_limits import MAX_MULTIPART_BYTES
    assert MultiPartParser.spool_max_size >= MAX_MULTIPART_BYTES
    # This exceeds the whole-body bound, including multipart overhead, before parsing.
    response = client.post("/v1/upload", files={"file": ("oversize.txt", b"x" * (MAX_MULTIPART_BYTES + 1))}, headers=KEY)
    assert response.status_code == 413


def test_no_findings_analysis_state(client):
    # A NO_FINDINGS specimen must have supported, unambiguous temporal syntax.
    # The old MSA's unspecified-day/receipt deadline now legitimately requires
    # review; its no-false-conflict and review behavior has separate regression tests.
    text = 'Section 1. Timing\nThe Effective Date is 2026-01-01.\nSection 2. Timing\nThe Expiration Date is 2026-12-31.'
    uploaded = client.post("/v1/upload", files={"file": ("synthetic-no-findings.txt", text.encode())}, headers=KEY)
    assert uploaded.status_code == 200
    response = client.post("/v1/analyze", json={"doc_id": uploaded.json()["doc_id"]}, headers=KEY)
    assert response.status_code == 200
    job_id = response.json()["job_id"]
    for _ in range(120):
        state = client.get(f"/v1/jobs/{job_id}", headers=KEY).json()
        if state["status"] in {"COMPLETED", "FAILED"}:
            break
        time.sleep(0.25)
    assert state["status"] == "COMPLETED", state
    result = client.get(f"/v1/jobs/{job_id}/results", headers=KEY).json()
    assert result["total_flags"] == 0
    assert result["metadata"]["analysis_state"] == "NO_FINDINGS"
    for text in extract_reports(result)[2:]:
        assert "NO_FINDINGS" in text and "No statistical anomalies" in text


def test_retention_cleans_only_owned_derived_cache(client, tmp_path, monkeypatch):
    from src.embeddings import cache as cache_module
    from src.api.security import get_safe_client_id
    import numpy as np
    cache_root = tmp_path / "vectors"
    monkeypatch.setattr(cache_module, "EMBEDDINGS_CACHE_DIR", cache_root)
    monkeypatch.setattr(storage, "EMBEDDINGS_CACHE_DIR", cache_root)
    doc_id = "doc_release_retention"
    storage.save_encrypted_document(doc_id, SAMPLE, get_safe_client_id(KEY["X-API-Key"]))
    clauses = get_ml_segmenter().segment(SAMPLE, doc_id)
    cache = cache_module.EmbeddingCache("release-test-model")
    cache.save_embeddings(clauses, np.ones((len(clauses), 3)))
    own = cache_root / "release-test-model" / f"{doc_id}.npz"
    other = own.with_name("doc_other.npz")
    other.write_bytes(own.read_bytes())
    assert own.exists()
    storage.delete_document_files(doc_id)
    assert not own.exists() and other.exists()
    assert not storage.document_exists(doc_id)
