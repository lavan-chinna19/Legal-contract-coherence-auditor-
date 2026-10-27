"""Real authenticated ingestion, saved evidence and local renderers; no model substitution."""
from io import BytesIO
import time
import pytest
from fastapi.testclient import TestClient
from src.api.main import app
from src.api import storage,jobs,rate_limiter,retention
from src.api.ingestion import parse_document

SOURCE=('Section 8. Payment\nSee Section 4, above.\n'
        'If price < 100, see Section 99; if units > 20, deliver.\n'
        'Section 4. Terms\nThe services are defined here.\n'
        'Section 4.2(a). Fees\nThe fee is agreed.\n'
        'Section 9. Reference\nSee Section 4.2(a).')
OWNER={'X-API-Key':'structural_test_owner'}
OTHER={'X-API-Key':'structural_test_other'}

@pytest.fixture
def client(tmp_path,monkeypatch):
    monkeypatch.setenv('API_KEYS','structural_test_owner,structural_test_other')
    monkeypatch.setattr(storage,'UPLOAD_DIR',tmp_path)
    monkeypatch.setattr(retention,'UPLOAD_DIR',tmp_path)
    monkeypatch.setattr(rate_limiter,'rate_limiter',rate_limiter.SlidingWindowRateLimiter())
    return TestClient(app)

def content(format):
    if format=='txt':return SOURCE.encode('utf8')
    out=BytesIO()
    if format=='docx':
        from docx import Document
        doc=Document()
        for line in SOURCE.splitlines():doc.add_paragraph(line)
        doc.save(out)
    else:
        from reportlab.pdfgen.canvas import Canvas
        canvas=Canvas(out)
        for i,line in enumerate(SOURCE.splitlines()):canvas.drawString(30,790-i*24,line)
        canvas.save()
    return out.getvalue()

@pytest.mark.parametrize('format',['txt','pdf','docx'])
def test_real_structural_release_workflow(client,format):
    data=content(format);extracted=parse_document(data,'source.'+format)
    uploaded=client.post('/v1/upload',headers=OWNER,files={'file':('source.'+format,data)})
    assert uploaded.status_code==200,uploaded.text
    doc_id=uploaded.json()['doc_id'];preview_url=f'/v1/documents/{doc_id}/clauses'
    assert client.get(preview_url).status_code==401
    assert client.get(preview_url,headers=OTHER).status_code==404
    preview=client.get(preview_url,headers=OWNER).json()
    assert preview['original_source_text']==extracted
    structure=preview['structural_document']
    assert structure['coordinate_space']=='original_extracted_text'
    assert ''.join(u['text'] for u in structure['units'])==extracted
    assert {s['identity'] for s in structure['sections']}=={'8','4','4.2(a)','9'}
    assert any(r['target_id']=='99' and r['resolution_status']=='MISSING' for r in preview['reference_resolutions'])
    assert any(r['target_id']=='4.2(a)' and r['resolution_status']=='RESOLVED' for r in preview['reference_resolutions'])
    started=client.post('/v1/analyze',headers=OWNER,json={'doc_id':doc_id})
    assert started.status_code==200,started.text
    job_id=started.json()['job_id']
    for _ in range(240):
        status=client.get(f'/v1/jobs/{job_id}',headers=OWNER).json()
        if status['status'] in ['COMPLETED','FAILED']:break
        time.sleep(.25)
    assert status['status']=='COMPLETED',status
    result=client.get(f'/v1/jobs/{job_id}/results',headers=OWNER).json()
    assert result['metadata']['clause_inventory']==preview['clauses']
    assert result['metadata']['structural_document']==structure
    assert any(f['evidence'].get('target_id')=='99' for f in result['rule_flags'])
    assert not any(f['evidence'].get('target_id')=='4.2(a)' for f in result['rule_flags'])
    for flag in result['rule_flags']:
        ev=flag['evidence']
        assert extracted[ev['char_start']:ev['char_end']]==ev['citation_text']
    graph_url='/v1/experimental/dependencies';payload={'job_id':job_id,'experimental':True}
    assert client.post(graph_url,headers=OTHER,json=payload).status_code==404
    graph_response=client.post(graph_url,headers=OWNER,json=payload)
    assert graph_response.status_code==200,graph_response.text
    graph=graph_response.json();assert len(graph['findings'])==1
    for edge in graph['findings']:
        assert extracted[edge['source_char_start']:edge['source_char_end']]==edge['reference']
        target=edge['target'];assert extracted[target['source_char_start']:target['source_char_end']]==target['excerpt']
    jobs.JOB_STORE.pop(job_id)
    reopened=client.get(f'/v1/jobs/{job_id}/results',headers=OWNER).json()
    assert reopened['metadata']['structural_document']==structure
    assert reopened['metadata']['dependency_graph']==graph
    for export in ['pdf','docx']:
        url=f'/v1/jobs/{job_id}/report/{export}'
        assert client.get(url).status_code==401
        assert client.get(url,headers=OTHER).status_code==404
        response=client.get(url,headers=OWNER)
        assert response.status_code==200 and len(response.content)>1000
        if export=='pdf':
            import pdfplumber
            with pdfplumber.open(BytesIO(response.content)) as pdf:text='\n'.join(p.extract_text() or '' for p in pdf.pages)
        else:
            from docx import Document
            text='\n'.join(p.text for p in Document(BytesIO(response.content)).paragraphs)
        assert 'Source-preserving structural evidence' in text
        assert 'original_extracted_text' in text and 'Section 99' in text
        assert structure['source_sha256'] in text

def test_review_state_survives_without_missing_target_flags(client):
    source='Section 4. First\nSection 4. Second\nSection 9. Reference\nSee Section 4.'
    uploaded=client.post('/v1/upload',headers=OWNER,files={'file':('ambiguous.txt',source.encode())}).json()
    started=client.post('/v1/analyze',headers=OWNER,json={'doc_id':uploaded['doc_id']}).json()
    for _ in range(240):
        status=client.get(f"/v1/jobs/{started['job_id']}",headers=OWNER).json()
        if status['status'] in ['COMPLETED','FAILED']:break
        time.sleep(.25)
    assert status['status']=='COMPLETED',status
    result=client.get(f"/v1/jobs/{started['job_id']}/results",headers=OWNER).json()
    assert result['metadata']['reference_review_count']>0
    assert result['metadata']['analysis_state']=='REVIEW_REQUIRED'
    assert not any(f['flag_type'].startswith('DANGLING_') for f in result['rule_flags'])
