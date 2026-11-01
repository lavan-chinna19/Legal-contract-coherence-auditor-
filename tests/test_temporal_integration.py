"""Actual local model, HTTP ingestion, persisted evidence and both export renderers."""
from io import BytesIO
import time
import pytest
from fastapi.testclient import TestClient
from src.api.main import app
from src.api import storage,jobs,retention,rate_limiter
from src.api.ingestion import parse_document

OWNER={'X-API-Key':'temporal_test_owner'}
OTHER={'X-API-Key':'temporal_test_other'}
SOURCE=('Section 1. Invoice\nInvoice Alpha is received on 2026-01-01.\n'
        'Section 2. Payment\nBuyer shall pay Invoice Alpha within 10 calendar days after receipt of Invoice Alpha.\n'
        'Section 3. Payment\nBuyer shall pay Invoice Alpha on 2026-02-01.')

@pytest.fixture
def client(tmp_path,monkeypatch):
    monkeypatch.setenv('API_KEYS','temporal_test_owner,temporal_test_other')
    monkeypatch.setattr(storage,'UPLOAD_DIR',tmp_path)
    monkeypatch.setattr(retention,'UPLOAD_DIR',tmp_path)
    monkeypatch.setattr(rate_limiter,'rate_limiter',rate_limiter.SlidingWindowRateLimiter())
    return TestClient(app)

def encoded(text,fmt):
    if fmt=='txt':return text.encode('utf8')
    out=BytesIO()
    if fmt=='docx':
        from docx import Document
        doc=Document()
        for line in text.splitlines():doc.add_paragraph(line)
        doc.save(out)
    else:
        from reportlab.pdfgen.canvas import Canvas
        canvas=Canvas(out,pagesize=(1200,900))
        for i,line in enumerate(text.splitlines()):canvas.drawString(20,870-i*22,line)
        canvas.save()
    return out.getvalue()

def analyze(client,text,fmt):
    content=encoded(text,fmt);source=parse_document(content,'temporal.'+fmt)
    uploaded=client.post('/v1/upload',headers=OWNER,files={'file':('temporal.'+fmt,content)})
    assert uploaded.status_code==200,uploaded.text
    doc_id=uploaded.json()['doc_id']
    started=client.post('/v1/analyze',headers=OWNER,json={'doc_id':doc_id})
    assert started.status_code==200,started.text
    job_id=started.json()['job_id']
    for _ in range(240):
        state=client.get(f'/v1/jobs/{job_id}',headers=OWNER).json()
        if state['status'] in ['COMPLETED','FAILED']:break
        time.sleep(.25)
    assert state['status']=='COMPLETED',state
    result=client.get(f'/v1/jobs/{job_id}/results',headers=OWNER).json()
    return doc_id,job_id,source,result

@pytest.mark.parametrize('fmt',['txt','pdf','docx'])
def test_temporal_auth_history_exports(client,fmt):
    doc_id,job_id,source,result=analyze(client,SOURCE,fmt)
    meta=result['metadata']['temporal_analysis']
    assert meta['confirmed_conflict_count']==1
    assert meta['coordinate_space']=='original_extracted_text'
    assert result['metadata']['analysis_state']=='REVIEW_REQUIRED'
    for s in meta['statements']:
        assert source[s['char_start']:s['char_end']]==s['raw_text']
        assert s['section_uid'] and s['section_identity']
    for suffix in ['/results','/report/pdf','/report/docx']:
        url=f'/v1/jobs/{job_id}'+suffix
        assert client.get(url).status_code==401
        assert client.get(url,headers=OTHER).status_code==404
    jobs.JOB_STORE.pop(job_id)
    restored=client.get(f'/v1/jobs/{job_id}/results',headers=OWNER).json()
    assert restored['metadata']['temporal_analysis']==meta
    assert client.get('/v1/documents',headers=OWNER).json()['documents'][0]['job_id']==job_id
    for export in ['pdf','docx']:
        response=client.get(f'/v1/jobs/{job_id}/report/{export}',headers=OWNER)
        assert response.status_code==200 and len(response.content)>1000
        assert response.headers['cache-control']=='no-store'
        if export=='pdf':
            import pdfplumber
            with pdfplumber.open(BytesIO(response.content)) as pdf:text='\n'.join(p.extract_text() or '' for p in pdf.pages)
        else:
            from docx import Document
            text='\n'.join(p.text for p in Document(BytesIO(response.content)).paragraphs)
        for expected in ['Temporal analysis','CONFIRMED_TEMPORAL_CONFLICT','RELATIVE_DEADLINE','2026-02-01','calendar','within']:
            assert expected in text
        assert meta['source_sha256'] in text

def test_uncertain_deadline_persists_without_false_conflict(client):
    _,job_id,_,result=analyze(client,'Section 1. Payment\nBuyer shall pay Invoice Alpha within 5 business days after receipt of Invoice Alpha.','txt')
    assert result['metadata']['temporal_review_count']>0
    assert not result['metadata']['temporal_analysis']['confirmed_conflict_count']
    assert result['metadata']['analysis_state']=='REVIEW_REQUIRED'
    jobs.JOB_STORE.pop(job_id)
    restored=client.get(f'/v1/jobs/{job_id}/results',headers=OWNER).json()
    assert restored['metadata']['temporal_analysis']==result['metadata']['temporal_analysis']

def test_original_msa_now_requires_review_without_false_alert(client):
    from src.redteam.fixtures import SYNTHETIC_FIXTURE_REGISTRY
    text='\n'.join(c.text for c in SYNTHETIC_FIXTURE_REGISTRY['SYNTH_DOC_01_MSA']())
    _,_,_,result=analyze(client,text,'txt')
    assert result['total_flags']==0
    assert result['metadata']['analysis_state']=='REVIEW_REQUIRED'
    assert result['metadata']['temporal_review_count']>0
    assert result['metadata']['temporal_analysis']['confirmed_conflict_count']==0
