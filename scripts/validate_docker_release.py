"""Real local Compose acceptance using synthetic documents; never removes volumes."""
import hashlib
import io
import json
import re
import subprocess
import time
from pathlib import Path

import httpx
import pdfplumber
from docx import Document
from dotenv import dotenv_values
from reportlab.pdfgen import canvas

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "tmp/docker-validation"
ENV_FILE = ROOT / ".env.docker.local"
COMPOSE = ["docker", "compose", "--env-file", str(ENV_FILE), "-p", "lexaudit-release"]
TEXT = "\n\n".join([
    "Section 3. Payment\nPayment for Services defined in Section 1 above is due within thirty days.",
    "Section 1. Definitions\nServices means the software work described in this agreement.",
    "Section 2. Confidentiality\nEach party shall protect confidential information disclosed by the other party.",
    "Section 4. Termination\nEither party may terminate this agreement upon thirty days written notice.",
])


class Validation:
    def __init__(self):
        OUT.mkdir(parents=True, exist_ok=True)
        self.config = dotenv_values(ENV_FILE)
        self.secrets = [v for k, v in self.config.items() if v and ("KEY" in k or "PASSWORD" in k)]
        self.secrets.extend((self.config.get("API_KEYS") or "").split(","))
        self.backend = f"http://127.0.0.1:{self.config.get('LEXAUDIT_BACKEND_PORT', '8011')}"
        self.frontend = f"http://127.0.0.1:{self.config.get('LEXAUDIT_FRONTEND_PORT', '3001')}"
        self.client = httpx.Client(base_url=self.backend, headers={"X-API-Key": self.config["API_KEY"]}, timeout=120)
        self.data = {"status": "IN_PROGRESS", "commands": [], "checks": []}

    def redact(self, text):
        for secret in sorted(self.secrets, key=len, reverse=True):
            if secret:
                text = text.replace(secret, "[REDACTED]")
        return text

    def save(self):
        (OUT / "acceptance.json").write_text(self.redact(json.dumps(self.data, indent=2)), encoding="utf-8")

    def command(self, args, name, timeout=120):
        started = time.perf_counter()
        result = subprocess.run(args, cwd=ROOT, capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=timeout)
        output = self.redact(result.stdout + result.stderr)
        (OUT / f"{name}.log").write_text(output, encoding="utf-8")
        self.data["commands"].append({"name": name, "argv": args, "exit_code": result.returncode,
                                      "seconds": round(time.perf_counter() - started, 3), "log": f"{name}.log"})
        self.save()
        print(f"COMMAND {name}: exit {result.returncode}", flush=True)
        if result.returncode:
            raise RuntimeError(f"{name} failed: {output[-2000:]}")
        return result.stdout

    def check(self, name, condition, detail=None):
        self.data["checks"].append({"name": name, "status": "PASS" if condition else "FAIL", "detail": detail})
        self.save()
        print(f"{'PASS' if condition else 'FAIL'} {name}", flush=True)
        if not condition:
            raise AssertionError(name)

    def request(self, name, method, route, expected=200, **kwargs):
        response = self.client.request(method, route, **kwargs)
        self.check(name, response.status_code == expected, {"http_status": response.status_code})
        return response

    def states(self, name):
        output = self.command(COMPOSE + ["ps", "--format", "json"], name)
        try:
            decoded = json.loads(output)
            items = decoded if isinstance(decoded, list) else [decoded]
        except json.JSONDecodeError:
            items = [json.loads(line) for line in output.splitlines() if line.strip()]
        self.check(name + " all four services healthy", len(items) == 4 and all(item.get("Health") == "healthy" for item in items))
        return {item["Service"]: item["ID"] for item in items}

    def wait_for_services(self, name):
        # An existing gateway can be temporarily unhealthy after planned backend
        # maintenance. Wait for recovery rather than treating old health as final.
        for attempt in range(60):
            output = self.command(COMPOSE + ["ps", "--format", "json"], f"{name}-{attempt}")
            try:
                decoded = json.loads(output)
                items = decoded if isinstance(decoded, list) else [decoded]
            except json.JSONDecodeError:
                items = [json.loads(line) for line in output.splitlines() if line.strip()]
            if len(items) == 4 and all(item.get("Health") == "healthy" for item in items):
                return
            time.sleep(3)
        raise TimeoutError("Required containers did not recover to healthy within the bounded wait")

    def export(self, job_id, doc_id, name):
        for fmt in ("pdf", "docx"):
            response = self.request(f"{name} {fmt} export", "GET", f"/v1/jobs/{job_id}/report/{fmt}")
            (OUT / f"{name}-report.{fmt}").write_bytes(response.content)
            if fmt == "pdf":
                with pdfplumber.open(io.BytesIO(response.content)) as pdf:
                    content = "\n".join(page.extract_text() or "" for page in pdf.pages)
            else:
                document = Document(io.BytesIO(response.content))
                content = "\n".join(p.text for p in document.paragraphs)
            self.check(f"{name} {fmt} valid factual content", doc_id in content and "LexAudit" in content and "Limitations" in content and "Payment" in content)
            self.check(f"{name} {fmt} protected download headers", response.headers.get("cache-control") == "no-store" and "attachment" in response.headers.get("content-disposition", ""))

    def run(self):
        self.command(COMPOSE + ["config", "--quiet"], "compose-config")
        self.command(COMPOSE + ["--profile", "tools", "run", "--rm", "--no-deps", "model-setup"], "seed-public-model-cache", timeout=600)
        self.command(COMPOSE + ["up", "-d"], "compose-start", timeout=420)
        self.wait_for_services("startup-health")
        self.command(COMPOSE + ["up", "-d", "--wait", "--wait-timeout", "300"], "compose-up", timeout=420)
        before_ids = self.states("initial-services")
        self.request("model readiness", "GET", "/ready")
        self.request("unauthenticated history denied", "GET", "/v1/documents", expected=401, headers={"X-API-Key": ""})
        self.request("invalid credential denied", "GET", "/v1/documents", expected=401, headers={"X-API-Key": "invalid-synthetic-key"})
        self.request("owner authenticated", "GET", "/v1/session")
        page = httpx.get(self.frontend, timeout=30)
        self.check("container frontend HTTP", page.status_code == 200 and "LexAudit" in page.text)
        paths = set(re.findall(r'<script[^>]+src="([^"]+)"', page.text))
        bundles = b"".join(httpx.get(self.frontend + p, timeout=30).content for p in paths if p.startswith("/"))
        self.check("frontend API URL baked at build time", self.backend.encode() in bundles)
        self.check("frontend has no private configuration secrets", all(v.encode() not in bundles for v in self.secrets if v))
        snapshots = []
        for fmt in ("txt", "pdf", "docx"):
            payload = io.BytesIO()
            if fmt == "txt":
                payload.write(TEXT.encode())
            elif fmt == "pdf":
                pdf = canvas.Canvas(payload)
                y = 780
                for line in TEXT.splitlines():
                    pdf.drawString(36, y, line)
                    y -= 18
                pdf.save()
            else:
                document = Document()
                for line in TEXT.splitlines():
                    document.add_paragraph(line)
                document.save(payload)
            doc_id = self.request(f"{fmt} upload", "POST", "/v1/upload", files={"file": (f"docker-synthetic.{fmt}", payload.getvalue())}).json()["doc_id"]
            parsed = self.request(f"{fmt} clause preview", "GET", f"/v1/documents/{doc_id}/clauses").json()
            job_id = self.request(f"{fmt} analysis start", "POST", "/v1/analyze", json={"doc_id": doc_id}).json()["job_id"]
            for _ in range(180):
                status = self.client.get(f"/v1/jobs/{job_id}")
                status.raise_for_status()
                state = status.json()
                if state["status"] == "COMPLETED":
                    break
                if state["status"] == "FAILED":
                    raise RuntimeError(state.get("error", "Model job failed"))
                time.sleep(1)
            else:
                raise TimeoutError("Analysis exceeded bounded polling")
            result = self.request(f"{fmt} actual inference results", "GET", f"/v1/jobs/{job_id}/results").json()
            self.check(f"{fmt} preview/result identity", parsed["clauses"] == result["metadata"]["clause_inventory"] and result["total_clauses"] == 4)
            self.check(f"{fmt} source evidence integrity", all(result["metadata"]["source_text"][c["char_start"]:c["char_end"]] == c["text"] for c in parsed["clauses"]))
            graph = self.request(f"{fmt} graph", "POST", "/v1/experimental/dependencies", json={"job_id": job_id, "experimental": True}).json()
            self.check(f"{fmt} directional graph evidence", any(f["status"] == "VERIFIED_TEXTUAL_CONFLICT" for f in graph["findings"]))
            if fmt == "txt":
                target = result["metadata"]["clause_inventory"][0]["clause_id"]
                params = {"doc_id": doc_id, "clause_id": target}
                attribution = self.request("Docker attribution", "GET", "/v1/xai/attribution", params=params).json()
                self.check("Docker attribution uses real tokens", bool(attribution["ig_payload"]["tokens"]) and attribution["clause_id"] == target)
                sensitivity = self.request("Docker neighbor sensitivity", "GET", "/v1/xai/sensitivity", params=params).json()
                self.check("Docker sensitivity target identity", sensitivity["clause_id"] == target)
                repair = self.request("Docker counterfactual", "POST", "/v1/experimental/reposition", json={**params, "experimental": True}).json()
                self.check("Docker counterfactual with abstention state", repair["target_clause_id"] == target and "is_abstaining" in repair)
                robust = self.request("Docker robustness negative control", "POST", "/v1/experimental/redteam", json={"doc_id": doc_id, "perturbation_type": "P9_NEGATIVE_CONTROL", "experimental": True}).json()
                self.check("Docker negative-control invariant", robust["original_checksum"] == robust["perturbed_checksum"])
            self.export(job_id, doc_id, fmt)
            other = (self.config["API_KEYS"] or "").split(",")[1]
            self.request(f"{fmt} cross-owner export denied", "GET", f"/v1/jobs/{job_id}/report/pdf", expected=404, headers={"X-API-Key": other})
            self.request(f"{fmt} cross-owner preview denied", "GET", f"/v1/documents/{doc_id}/clauses", expected=404, headers={"X-API-Key": other})
            latest = self.request(f"{fmt} saved snapshot", "GET", f"/v1/jobs/{job_id}/results").json()
            snapshots.append({"doc_id": doc_id, "job_id": job_id, "digest": hashlib.sha256(json.dumps(latest, sort_keys=True).encode()).hexdigest()})
        history = self.request("saved owner history", "GET", "/v1/documents").json()["documents"]
        self.check("all three analyses in history", all(any(d["doc_id"] == x["doc_id"] and d["job_id"] == x["job_id"] for d in history) for x in snapshots))
        sample = snapshots[0]
        saved = self.client.get(f"/v1/jobs/{sample['job_id']}/results").json()
        clause = saved["metadata"]["clause_inventory"][0]["clause_id"]
        feedback = self.request("PostgreSQL-backed synthetic feedback", "POST", "/v1/feedback", json={"doc_id": sample["doc_id"], "clause_id": clause, "original_severity": "MEDIUM", "reviewer_verdict": "REJECT_ANOMALY", "reviewer_id": "docker-release-synthetic", "provenance": "SYNTHETIC_TEST"}).json()
        self.data["feedback_id"] = feedback["feedback_id"]
        self.data["snapshots"] = snapshots
        self.save()
        runtime = '''import os,json,socket,torch
from pathlib import Path
from src.feedback.storage import get_all_feedback,is_postgres
from src.api.storage import load_and_decrypt_document
checks={"uid":os.getuid(),"torch":torch.__version__,"cuda":torch.cuda.is_available(),"offline":os.environ.get("HF_HUB_OFFLINE"),"download":os.environ.get("LEXAUDIT_ALLOW_MODEL_DOWNLOAD"),"postgres":is_postgres(),"synthetic_feedback":any(x.reviewer_id=="docker-release-synthetic" for x in get_all_feedback())}
try:
 socket.create_connection(("1.1.1.1",443),timeout=2).close(); checks["internet_egress_blocked"]=False
except OSError: checks["internet_egress_blocked"]=True
print(json.dumps(checks))'''
        details = json.loads(self.command(COMPOSE + ["exec", "-T", "backend", "python", "-c", runtime], "runtime-locality"))
        self.check("local CPU offline inference configuration", details["uid"] == 10001 and "+cpu" in details["torch"] and not details["cuda"] and details["offline"] == "1" and details["download"] == "0", details)
        self.check("runtime external network is isolated", details["internet_egress_blocked"])
        self.check("feedback actually persisted in PostgreSQL", details["postgres"] and details["synthetic_feedback"])
        backend_id = self.command(COMPOSE + ["ps", "-q", "backend"], "backend-id").strip()
        mount_info = json.loads(self.command(["docker", "inspect", "--format", "{{json .Mounts}}", backend_id], "backend-mounts"))
        self.check("model is a read-only local volume", any(m["Type"] == "volume" and not m["RW"] and m["Destination"] == "/opt/hf/hub" for m in mount_info))
        self.check("persistent uploads and vectors mounted", all(any(m["Type"] == "volume" and m["Destination"] == p for m in mount_info) for p in ["/app/data/uploads", "/app/data/processed/embeddings_cache", "/app/data/feedback"]))
        self.command(COMPOSE + ["logs", "--no-color", "--tail", "150", "backend", "frontend", "postgres"], "pre-restart-service-logs")
        self.command(COMPOSE + ["restart", "postgres", "backend", "frontend"], "compose-restart", timeout=150)
        self.wait_for_services("restart-health")
        self.command(COMPOSE + ["up", "-d", "--wait", "--wait-timeout", "300"], "wait-after-restart", timeout=330)
        self.states("services-after-restart")
        for item in snapshots:
            response = self.request("analysis survives restart " + item["doc_id"], "GET", f"/v1/jobs/{item['job_id']}/results").json()
            digest = hashlib.sha256(json.dumps(response, sort_keys=True).encode()).hexdigest()
            self.check("exact saved scores/evidence after restart " + item["doc_id"], digest == item["digest"])
        self.export(sample["job_id"], sample["doc_id"], "after-restart")
        self.command(COMPOSE + ["logs", "--no-color", "--tail", "100", "backend", "frontend", "postgres"], "after-restart-service-logs")
        self.command(COMPOSE + ["up", "-d", "--force-recreate", "--wait", "--wait-timeout", "300", "backend", "frontend"], "compose-recreate", timeout=360)
        after_ids = self.states("services-after-recreate")
        self.check("new application containers created", before_ids["backend"] != after_ids["backend"] and before_ids["frontend"] != after_ids["frontend"])
        for item in snapshots:
            response = self.request("analysis survives recreation " + item["doc_id"], "GET", f"/v1/jobs/{item['job_id']}/results").json()
            self.check("exact persisted snapshot after recreation " + item["doc_id"], hashlib.sha256(json.dumps(response, sort_keys=True).encode()).hexdigest() == item["digest"])
        final_details = json.loads(self.command(COMPOSE + ["exec", "-T", "backend", "python", "-c", runtime], "runtime-after-recreate"))
        self.check("feedback survives database restart and app recreation", final_details["synthetic_feedback"])
        self.export(sample["job_id"], sample["doc_id"], "after-recreate")
        self.command(COMPOSE + ["logs", "--no-color", "--tail", "150", "backend", "frontend", "postgres"], "final-service-logs")
        self.data["status"] = "PASS"
        self.save()
        print(f"PASS: {len(self.data['checks'])} real Docker checks; restart and recreation verified.", flush=True)


if __name__ == "__main__":
    validation = Validation()
    try:
        validation.run()
    except Exception as exc:
        validation.data["status"] = "FAIL"
        validation.data["error"] = validation.redact(str(exc))
        validation.save()
        raise
